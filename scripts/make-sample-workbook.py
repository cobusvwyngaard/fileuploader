"""Generates docs/sample/UploadsReport.xlsx: a workbook with the named table
"Uploads" containing every standard column plus table_number and description.

    python3 scripts/make-sample-workbook.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

STANDARD = [
    "submission_id", "upload_session_id", "session_title", "submitted_at_utc",
    "original_filename", "stored_filename", "file_size_bytes", "content_type",
    "onedrive_folder_path", "onedrive_web_url", "processing_status",
]
FIELDS = ["table_number", "description"]
WIDTHS = {"session_title": 28, "original_filename": 30, "stored_filename": 44, "onedrive_folder_path": 26, "onedrive_web_url": 50, "description": 40, "submitted_at_utc": 22}

wb = Workbook()
ws = wb.active
ws.title = "Uploads"
headers = STANDARD + FIELDS
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)
# One placeholder data row: Excel tables need at least one body row. Delete it
# in Excel after the first real submission arrives, or leave it.
ws.append(["" for _ in headers])
for i, h in enumerate(headers, start=1):
    ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTHS.get(h, max(14, len(h) + 2))
last_col = ws.cell(row=1, column=len(headers)).column_letter
table = Table(displayName="Uploads", ref=f"A1:{last_col}2")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(table)
ws.freeze_panes = "A2"

out = Path(__file__).resolve().parent.parent / "docs" / "sample" / "UploadsReport.xlsx"
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Wrote {out}")
